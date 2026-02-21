#!/usr/bin/env python3
"""
Desvio médio por UC + Concessionária (consumo de energia).
CLI utility consumed by ocr_service/app.py (/ocr/desvio-media).

v5 (upgrade):
- Baseline rolling/expanding por Mes_Ref (opcional) + lookback
- Normalização por dias faturados (kWh/dia) com fallback auto/on/off
- MAD robusto com escala 1.4826 + piso de dispersão para evitar score absurdo quando MAD=0
- Campos explicáveis por registro: baseline_center, baseline_disp, baseline_len, limites, valor_dia, baseline_center_dia...
- Detecção de salto mês-a-mês (SALTO_ALTO/SALTO_BAIXO)
- Impacto financeiro opcional via tarifa_estimada ou tarifa_col
- Série (include_series) para sparklines/gráficos
- Mantém estrutura de saída: summary, groups, anomalias, duplicidades
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from statistics import mean, median
from typing import Any, Dict, List, Optional, Tuple

ROBUST_MAD_SCALE = 1.4826  # torna MAD comparável ao desvio padrão (aprox, normal)


def norm_text(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip().lower()


def parse_float(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        x = float(v)
        if math.isnan(x):
            return None
        return x
    s = str(v).strip()
    if not s:
        return None
    s = s.replace("R$", "").replace(" ", "")
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        x = float(s)
        if math.isnan(x):
            return None
        return x
    except Exception:
        return None


def to_int(v: Any, default: int) -> int:
    try:
        return int(str(v).strip())
    except Exception:
        return default


def to_float(v: Any, default: float) -> float:
    try:
        return float(str(v).strip())
    except Exception:
        return default


def parse_placeholder_list(raw: str) -> set:
    out = set()
    if not raw:
        return out
    for p in str(raw).split(","):
        f = parse_float(p)
        if f is not None:
            out.add(float(f))
    return out


def coalesce_key(d: Dict[str, Any], key: str) -> Any:
    if key in d:
        return d.get(key)
    lk = key.lower()
    for k, v in d.items():
        if str(k).lower() == lk:
            return v
    return None


def parse_mes_ref(v: Any) -> str:
    """
    Retorna ISO YYYY-MM-DD (se possível), caso contrário devolve string original.
    Se vier mês/ano, força day=1.
    """
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    if not s:
        return ""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%Y", "%Y/%m", "%Y-%m"):
        try:
            dt = datetime.strptime(s, fmt)
            if fmt in ("%m/%Y", "%Y/%m", "%Y-%m"):
                dt = dt.replace(day=1)
            return dt.date().isoformat()
        except Exception:
            pass
    return s


def mesref_to_date(mes_ref: str) -> Optional[date]:
    if not mes_ref:
        return None
    try:
        return datetime.strptime(mes_ref[:10], "%Y-%m-%d").date()
    except Exception:
        return None


def load_rows(path: str, sheet: Optional[str], header_row: int) -> List[Dict[str, Any]]:
    p = Path(path)
    ext = p.suffix.lower()

    if ext in (".csv", ".txt"):
        data = p.read_bytes()
        text = None
        for enc in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except Exception:
                continue
        if text is None:
            raise RuntimeError("Nao foi possivel decodificar CSV")
        lines = text.splitlines()
        if not lines:
            return []
        reader = csv.reader(lines)
        all_rows = list(reader)
        if not all_rows:
            return []
        idx = max(0, header_row - 1)
        if idx >= len(all_rows):
            return []
        header = [str(h).strip() for h in all_rows[idx]]
        out = []
        for row in all_rows[idx + 1 :]:
            if not any(str(c).strip() for c in row):
                continue
            item = {}
            for i, h in enumerate(header):
                if not h:
                    continue
                item[h] = row[i] if i < len(row) else ""
            out.append(item)
        return out

    try:
        import openpyxl  # type: ignore
    except Exception as exc:
        raise RuntimeError("openpyxl nao instalado. Instale com: pip install openpyxl") from exc

    wb = openpyxl.load_workbook(filename=path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]

    rows_iter = ws.iter_rows(values_only=True)
    all_rows = list(rows_iter)
    if not all_rows:
        return []

    idx = max(0, header_row - 1)
    if idx >= len(all_rows):
        return []

    header = [str(h).strip() if h is not None else "" for h in all_rows[idx]]
    out = []
    for row in all_rows[idx + 1 :]:
        if not row or not any(c is not None and str(c).strip() for c in row):
            continue
        item = {}
        for i, h in enumerate(header):
            if not h:
                continue
            item[h] = row[i] if i < len(row) else None
        out.append(item)
    return out


# =========================
# Core math
# =========================

def _quantile(sorted_vals: List[float], q: float) -> float:
    if not sorted_vals:
        return 0.0
    n = len(sorted_vals)
    if n == 1:
        return float(sorted_vals[0])
    pos = (n - 1) * q
    lo = int(math.floor(pos))
    hi = int(math.ceil(pos))
    if lo == hi:
        return float(sorted_vals[lo])
    w = pos - lo
    return float(sorted_vals[lo] * (1 - w) + sorted_vals[hi] * w)


def calc_baseline(values: List[float], method: str, min_disp: float, disp_floor_pct: float) -> Tuple[float, float]:
    """Retorna (center, dispersion). Dispersion sempre >= piso."""
    vals = [float(x) for x in values if x is not None and not math.isnan(float(x))]
    if not vals:
        return 0.0, max(1.0, min_disp)

    method = (method or "mean_mad").strip().lower()
    vals_sorted = sorted(vals)

    def apply_floor(center: float, disp: float) -> float:
        dyn = abs(center) * max(0.0, disp_floor_pct)
        floor = max(min_disp, 1e-9, dyn)
        return float(disp if disp > floor else floor)

    if method == "robust_mad":
        c = float(median(vals_sorted))
        dev = [abs(x - c) for x in vals_sorted]
        mad = float(median(dev)) if dev else 0.0
        disp = mad * ROBUST_MAD_SCALE
        if disp <= 0:
            q25 = _quantile(vals_sorted, 0.25)
            q75 = _quantile(vals_sorted, 0.75)
            iqr = float(q75 - q25)
            disp = (iqr / 1.349) if iqr > 0 else 0.0
        disp = apply_floor(c, disp)
        return c, disp

    if method == "zscore":
        c = float(mean(vals_sorted))
        if len(vals_sorted) >= 2:
            try:
                import statistics
                disp = float(statistics.pstdev(vals_sorted))
            except Exception:
                disp = 0.0
        else:
            disp = 0.0
        disp = apply_floor(c, disp)
        return c, disp

    if method == "iqr":
        c = float(median(vals_sorted))
        q25 = _quantile(vals_sorted, 0.25)
        q75 = _quantile(vals_sorted, 0.75)
        iqr = float(q75 - q25)
        disp = (iqr / 1.349) if iqr > 0 else 0.0
        disp = apply_floor(c, disp)
        return c, disp

    c = float(mean(vals_sorted))
    disp = float(mean([abs(x - c) for x in vals_sorted])) if vals_sorted else 0.0
    disp = apply_floor(c, disp)
    return c, disp


# =========================
# Domain features
# =========================

def mark_placeholder_repeated(items: List[Dict[str, Any]], threshold: int) -> None:
    vals = [x.get("value") for x in items if x.get("status") == "PLACEHOLDER"]
    cnt = Counter(vals)
    repeated = {v for v, c in cnt.items() if c >= threshold}
    if not repeated:
        return
    for it in items:
        if it.get("status") == "PLACEHOLDER" and it.get("value") in repeated:
            it["status"] = "PLACEHOLDER_REPETIDO"


def decide_normalize_by_days(items: List[Dict[str, Any]], normalize_mode: str, coverage_threshold: float) -> bool:
    m = (normalize_mode or "auto").strip().lower()
    if m in ("0", "false", "off", "nao", "não"):
        return False
    if m in ("1", "true", "on", "sim", "yes"):
        return True

    valid = [x for x in items if x.get("status") not in ("INVALIDO", "PLACEHOLDER", "PLACEHOLDER_REPETIDO") and (x.get("value") or 0) > 0]
    if not valid:
        return False
    has_days = [x for x in valid if (x.get("days") or 0) > 0]
    cov = len(has_days) / max(1, len(valid))
    return cov >= max(0.0, min(1.0, coverage_threshold))


def get_history_slice(series: List[Dict[str, Any]], idx: int, baseline_mode: str, lookback: int) -> List[Dict[str, Any]]:
    mode = (baseline_mode or "all_valid").strip().lower()
    if mode in ("all_valid", "global", "full"):
        return series
    prev = series[:idx]
    if mode in ("expanding", "prev", "prev_only"):
        return prev
    if mode in ("rolling", "window"):
        if lookback and lookback > 0:
            return prev[-lookback:]
        return prev
    return series


def maybe_aggregate_duplicates(items: List[Dict[str, Any]], agg: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    agg = (agg or "none").strip().lower()
    if agg in ("none", "nao", "não", "off", "0", ""):
        return items, []
    if not any(it.get("mes_ref") for it in items):
        return items, []

    by_mes: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for it in items:
        by_mes[it.get("mes_ref") or ""].append(it)

    out: List[Dict[str, Any]] = []
    dups: List[Dict[str, Any]] = []

    for mr, lst in by_mes.items():
        if len(lst) == 1:
            out.append(lst[0])
            continue
        vals = [x.get("value") for x in lst if x.get("value") is not None]
        days = [x.get("days") for x in lst if (x.get("days") or 0) > 0]
        if not vals:
            out.append(lst[0])
            continue

        if agg == "sum":
            new_val = float(sum(float(v) for v in vals))
        elif agg == "mean":
            new_val = float(mean([float(v) for v in vals]))
        elif agg == "max":
            new_val = float(max(float(v) for v in vals))
        elif agg == "min":
            new_val = float(min(float(v) for v in vals))
        else:
            new_val = float(mean([float(v) for v in vals]))

        new_days = float(max(days)) if days else (float(lst[0].get("days") or 0) or 0.0)

        base = dict(lst[0])
        base["value"] = new_val
        base["days"] = new_days
        base["row_index"] = int(lst[0].get("row_index") or 0)
        base["status"] = "NORMAL" if base.get("status") != "INVALIDO" else base.get("status")
        out.append(base)

        dups.append(
            {
                "uc": base.get("uc", ""),
                "concessionaria": base.get("concessionaria", ""),
                "mes_ref": mr,
                "qtd": len(lst),
                "agg": agg,
                "media_kwh": float(mean([float(v) for v in vals])),
                "soma_kwh": float(sum(float(v) for v in vals)),
                "alerta": "DUPLICIDADE/AGREGACAO",
            }
        )

    out.sort(key=lambda x: (mesref_to_date(x.get("mes_ref", "") or "") or date.min, str(x.get("mes_ref", ""))))
    return out, dups


@dataclass
class Anomalia:
    uc: str
    concessionaria: str
    mes_ref: str
    valor: float
    row_index: int
    status: str
    status_extra: str
    score: float
    dif_abs: float
    dif_pct: Optional[float]
    normal_value: float

    dias_faturados: Optional[float] = None
    valor_dia: Optional[float] = None
    baseline_center: Optional[float] = None
    baseline_disp: Optional[float] = None
    baseline_len: Optional[int] = None
    limite_inf: Optional[float] = None
    limite_sup: Optional[float] = None
    baseline_center_dia: Optional[float] = None
    baseline_disp_dia: Optional[float] = None
    impacto_rs: Optional[float] = None
    impacto_abs_rs: Optional[float] = None
    jump_pct: Optional[float] = None


def analyze_group(
    items_raw: List[Dict[str, Any]],
    *,
    min_base: int,
    score_threshold: float,
    pct_high: float,
    pct_low: float,
    k: float,
    method: str,
    leitura_estimada_n: int,
    baseline_mode: str,
    lookback: int,
    include_series: bool,
    series_limit: int,
    normalize_by_days: str,
    days_coverage_threshold: float,
    min_disp: float,
    disp_floor_pct: float,
    jump_high: Optional[float],
    jump_low: Optional[float],
    jump_min_base: int,
    tarifa_estimada: Optional[float],
    duplicates_agg: str,
) -> Dict[str, Any]:
    items, dups_local = maybe_aggregate_duplicates(items_raw, duplicates_agg)

    def sort_key(it: Dict[str, Any]):
        d = mesref_to_date(it.get("mes_ref", "") or "")
        return (d or date.min, str(it.get("mes_ref", "")))

    items_sorted = sorted(items, key=sort_key)

    do_norm = decide_normalize_by_days(items_sorted, normalize_by_days, days_coverage_threshold)

    series: List[Dict[str, Any]] = []
    for it in items_sorted:
        st = it.get("status", "NORMAL")
        v = it.get("value")
        d = it.get("days") or 0
        v_use: Optional[float] = None
        if v is not None and st not in ("INVALIDO", "PLACEHOLDER", "PLACEHOLDER_REPETIDO") and float(v) > 0:
            if do_norm and d and float(d) > 0:
                v_use = float(v) / float(d)
            else:
                v_use = float(v)
        series.append({**it, "value_use": v_use})

    leitura_estimada_values = set()
    base_vals_raw = [float(x["value"]) for x in items_sorted
                     if x.get("status") not in ("INVALIDO", "PLACEHOLDER", "PLACEHOLDER_REPETIDO")
                     and (x.get("value") or 0) > 0]
    if leitura_estimada_n > 1 and base_vals_raw:
        freq = Counter(round(float(v), 6) for v in base_vals_raw)
        leitura_estimada_values = {v for v, c in freq.items() if c >= leitura_estimada_n}

    base_use_vals_global = [x["value_use"] for x in series if x.get("value_use") is not None]
    global_center_use, global_disp_use = calc_baseline(base_use_vals_global, method, min_disp=min_disp, disp_floor_pct=disp_floor_pct)
    global_len = len(base_use_vals_global)

    days_valid = [float(x.get("days") or 0) for x in items_sorted if (x.get("days") or 0) > 0]
    days_med = float(median(days_valid)) if days_valid else 0.0
    if do_norm and days_med > 0:
        global_center_kwh = global_center_use * days_med
        global_disp_kwh = global_disp_use * days_med
    else:
        global_center_kwh = global_center_use
        global_disp_kwh = global_disp_use

    anomalias: List[Anomalia] = []
    prev_use: Optional[float] = None

    for idx, it in enumerate(series):
        status = it.get("status", "NORMAL")
        status_extra = ""
        score = 0.0
        dif_abs = 0.0
        dif_pct: Optional[float] = None
        jump_pct: Optional[float] = None

        v_raw = it.get("value")
        v_use = it.get("value_use")
        days = float(it.get("days") or 0) if it.get("days") is not None else 0.0

        hist_items = get_history_slice(series, idx, baseline_mode, lookback)
        hist_use_vals = [x.get("value_use") for x in hist_items if x.get("value_use") is not None]

        if baseline_mode.strip().lower() in ("rolling", "expanding", "prev", "prev_only", "window") and len(hist_use_vals) == 0:
            center_use, disp_use, blen = global_center_use, global_disp_use, global_len
        else:
            center_use, disp_use = calc_baseline(hist_use_vals, method, min_disp=min_disp, disp_floor_pct=disp_floor_pct)
            blen = len(hist_use_vals)

        if do_norm and days > 0 and v_raw is not None:
            baseline_center_dia = center_use
            baseline_disp_dia = disp_use
            baseline_center = center_use * days
            baseline_disp = disp_use * days
        else:
            baseline_center_dia = None
            baseline_disp_dia = None
            baseline_center = center_use
            baseline_disp = disp_use

        lim_inf = baseline_center - (k * baseline_disp)
        lim_sup = baseline_center + (k * baseline_disp)

        tiny_sample = blen < min_base

        if v_raw is not None and status not in ("INVALIDO", "PLACEHOLDER", "PLACEHOLDER_REPETIDO"):
            v_kwh = float(v_raw)
            dif_abs = abs(v_kwh - baseline_center)
            score = dif_abs / (baseline_disp if baseline_disp else 1.0)
            dif_pct = (v_kwh - baseline_center) / baseline_center if baseline_center else None

            if prev_use is not None and v_use is not None and prev_use > 0:
                jump_pct = (float(v_use) - float(prev_use)) / float(prev_use)
                if (jump_high is not None) and (blen >= jump_min_base) and jump_pct >= float(jump_high):
                    status_extra = "SALTO_ALTO" if not status_extra else f"{status_extra}|SALTO_ALTO"
                if (jump_low is not None) and (blen >= jump_min_base) and jump_pct <= float(jump_low):
                    status_extra = "SALTO_BAIXO" if not status_extra else f"{status_extra}|SALTO_BAIXO"

            if not tiny_sample:
                if v_kwh < lim_inf:
                    status = "MUITO_BAIXO"
                elif v_kwh > lim_sup:
                    status = "PICO_OUTLIER"
                elif score >= score_threshold:
                    status = "OUTLIER_SCORE"
                else:
                    status = "NORMAL"
            else:
                status = "NORMAL"
                status_extra = "AMOSTRA_PEQUENA" if not status_extra else f"{status_extra}|AMOSTRA_PEQUENA"

            if dif_pct is not None:
                if dif_pct >= pct_high:
                    status_extra = "ACIMA_100PCT" if not status_extra else f"{status_extra}|ACIMA_100PCT"
                elif dif_pct <= pct_low:
                    status_extra = "ABAIXO_90PCT" if not status_extra else f"{status_extra}|ABAIXO_90PCT"

            if round(float(v_kwh), 6) in leitura_estimada_values:
                status_extra = "LEITURA_ESTIMADA" if not status_extra else f"{status_extra}|LEITURA_ESTIMADA"

            impacto = None
            impacto_abs = None
            tarifa = tarifa_estimada if tarifa_estimada is not None else (parse_float(it.get("tarifa")) if it.get("tarifa") is not None else None)
            if tarifa is not None and dif_pct is not None:
                impacto = (v_kwh - baseline_center) * float(tarifa)
                impacto_abs = abs(impacto)

            if status != "NORMAL" or status_extra:
                anomalias.append(
                    Anomalia(
                        uc=it.get("uc", ""),
                        concessionaria=it.get("concessionaria", ""),
                        mes_ref=it.get("mes_ref", ""),
                        valor=float(v_kwh),
                        row_index=int(it.get("row_index") or 0),
                        status=status,
                        status_extra=status_extra,
                        score=float(score),
                        dif_abs=float(dif_abs),
                        dif_pct=float(dif_pct) if dif_pct is not None else None,
                        normal_value=float(baseline_center),
                        dias_faturados=float(days) if days > 0 else None,
                        valor_dia=float(v_kwh / days) if (days and days > 0) else None,
                        baseline_center=float(baseline_center),
                        baseline_disp=float(baseline_disp),
                        baseline_len=int(blen),
                        limite_inf=float(lim_inf),
                        limite_sup=float(lim_sup),
                        baseline_center_dia=float(baseline_center_dia) if baseline_center_dia is not None else None,
                        baseline_disp_dia=float(baseline_disp_dia) if baseline_disp_dia is not None else None,
                        impacto_rs=float(impacto) if impacto is not None else None,
                        impacto_abs_rs=float(impacto_abs) if impacto_abs is not None else None,
                        jump_pct=float(jump_pct) if jump_pct is not None else None,
                    )
                )

        if v_use is not None and float(v_use) > 0:
            prev_use = float(v_use)

    anomalias.sort(key=lambda x: x.dif_abs, reverse=True)

    out_series = None
    if include_series:
        pts = []
        for it in series:
            mr = it.get("mes_ref", "")
            v_kwh = it.get("value")
            d = it.get("days") or None
            pts.append(
                {
                    "mes_ref": mr,
                    "value": float(v_kwh) if v_kwh is not None else None,
                    "days": float(d) if d not in (None, "") else None,
                    "value_dia": (float(v_kwh) / float(d)) if (v_kwh is not None and d and float(d) > 0) else None,
                }
            )
        if series_limit and series_limit > 0 and len(pts) > series_limit:
            pts = pts[-series_limit:]
        out_series = pts

    return {
        "normal_value": float(global_center_kwh),
        "media_base": float(global_center_kwh),
        "disp": float(global_disp_kwh),
        "baseline_center_use": float(global_center_use),
        "baseline_disp_use": float(global_disp_use),
        "do_normalize_by_days": bool(do_norm),
        "days_median": float(days_med) if days_med else None,
        "limite_inf": float(global_center_kwh - (k * global_disp_kwh)),
        "limite_sup": float(global_center_kwh + (k * global_disp_kwh)),
        "base_len": int(global_len),
        "total_registros": len(items_sorted),
        "anomalias": [a.__dict__ for a in anomalias],
        "series": out_series,
        "duplicidades_local": dups_local,
    }


def make_json_safe(obj: Any) -> Any:
    if isinstance(obj, dict):
        return {k: make_json_safe(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [make_json_safe(v) for v in obj]
    if isinstance(obj, (datetime, date)):
        return obj.isoformat()
    return obj


def main() -> None:
    parser = argparse.ArgumentParser(description="Analise de desvio medio por UC + Concessionaria")
    parser.add_argument("--input", required=True)
    parser.add_argument("--unit-col", default="UC")
    parser.add_argument("--concessionaria-col", default="Concessionaria")
    parser.add_argument("--value-col", default="KWH_FPonta")
    parser.add_argument("--days-col", default="Dias_Faturados")
    parser.add_argument("--tarifa-col", default="")  # opcional: tarifa por linha

    parser.add_argument("--placeholder", default="0,1,50")
    parser.add_argument("--min-base", default="8")
    parser.add_argument("--score-threshold", default="3.0")
    parser.add_argument("--pct-high", default="1.0")
    parser.add_argument("--pct-low", default="-0.9")
    parser.add_argument("--k", default="3.0")
    parser.add_argument("--placeholder-repeat-threshold", default="3")
    parser.add_argument("--max-anom", default="200")
    parser.add_argument("--header-row", default="1")
    parser.add_argument("--sheet", default="")
    parser.add_argument("--mes-col", default="")
    parser.add_argument("--method", default="mean_mad")

    parser.add_argument("--baseline-mode", default="all_valid")  # all_valid|rolling|expanding
    parser.add_argument("--lookback", default="12")

    parser.add_argument("--normalize-by-days", default="auto")  # auto|on|off
    parser.add_argument("--days-coverage-threshold", default="0.7")
    parser.add_argument("--min-disp", default="1.0")
    parser.add_argument("--disp-floor-pct", default="0.10")

    parser.add_argument("--leitura-estimada-n", default="3")
    parser.add_argument("--jump-high", default="0.8")
    parser.add_argument("--jump-low", default="-0.6")
    parser.add_argument("--jump-min-base", default="4")

    parser.add_argument("--include-series", default="false")
    parser.add_argument("--series-limit", default="24")
    parser.add_argument("--tarifa-estimada", default="")

    parser.add_argument("--duplicates-agg", default="none")

    parser.add_argument("--pretty", action="store_true")
    args = parser.parse_args()

    rows = load_rows(args.input, sheet=args.sheet or None, header_row=to_int(args.header_row, 1))

    placeholder_values = parse_placeholder_list(args.placeholder)
    min_base = to_int(args.min_base, 8)
    score_threshold = to_float(args.score_threshold, 3.0)
    pct_high = to_float(args.pct_high, 1.0)
    pct_low = to_float(args.pct_low, -0.9)
    k = to_float(args.k, 3.0)
    max_anom = to_int(args.max_anom, 200)
    repeat_thr = to_int(args.placeholder_repeat_threshold, 3)
    leitura_estimada_n = to_int(args.leitura_estimada_n, 3)
    method = (args.method or "mean_mad").strip().lower()

    baseline_mode = (args.baseline_mode or "all_valid").strip().lower()
    lookback = to_int(args.lookback, 12)

    normalize_by_days = (args.normalize_by_days or "auto").strip().lower()
    days_cov = to_float(args.days_coverage_threshold, 0.7)
    min_disp = to_float(args.min_disp, 1.0)
    disp_floor_pct = to_float(args.disp_floor_pct, 0.10)

    jump_high = parse_float(args.jump_high)
    jump_low = parse_float(args.jump_low)
    jump_min_base = to_int(args.jump_min_base, 4)

    include_series = norm_text(args.include_series) in ("1", "true", "sim", "yes")
    series_limit = to_int(args.series_limit, 24)

    tarifa_estimada = parse_float(args.tarifa_estimada)
    tarifa_col = (args.tarifa_col or "").strip()
    duplicates_agg = (args.duplicates_agg or "none").strip().lower()

    mes_col = args.mes_col.strip() if args.mes_col else ""
    if not mes_col:
        for guess in ("Mes_Ref", "Mês_Ref", "MesRef", "Referencia", "Referência"):
            if rows and coalesce_key(rows[0], guess) is not None:
                mes_col = guess
                break

    prepared: List[Dict[str, Any]] = []
    grouped_raw: Dict[Tuple[str, str], List[Dict[str, Any]]] = defaultdict(list)

    for i, r in enumerate(rows, start=1):
        uc = str(coalesce_key(r, args.unit_col) or "").strip()
        conc = str(coalesce_key(r, args.concessionaria_col) or "").strip()
        v = parse_float(coalesce_key(r, args.value_col))
        days = parse_float(coalesce_key(r, args.days_col))
        mes_ref = parse_mes_ref(coalesce_key(r, mes_col)) if mes_col else ""

        status = "NORMAL"
        if v is None:
            status = "INVALIDO"
        elif float(v) in placeholder_values:
            status = "PLACEHOLDER"

        item = {"uc": uc, "concessionaria": conc, "mes_ref": mes_ref, "value": v, "days": days, "status": status, "row_index": i}
        if tarifa_col:
            item["tarifa"] = parse_float(coalesce_key(r, tarifa_col))
        prepared.append(item)
        grouped_raw[(uc, conc)].append(item)

    for _, gitems in grouped_raw.items():
        mark_placeholder_repeated(gitems, threshold=repeat_thr)

    groups: List[Dict[str, Any]] = []
    anomalias_flat: List[Dict[str, Any]] = []
    duplicidades_global: List[Dict[str, Any]] = []

    for (uc, conc), items in grouped_raw.items():
        stat = analyze_group(
            items,
            min_base=min_base,
            score_threshold=score_threshold,
            pct_high=pct_high,
            pct_low=pct_low,
            k=k,
            method=method,
            leitura_estimada_n=leitura_estimada_n,
            baseline_mode=baseline_mode,
            lookback=lookback,
            include_series=include_series,
            series_limit=series_limit,
            normalize_by_days=normalize_by_days,
            days_coverage_threshold=days_cov,
            min_disp=min_disp,
            disp_floor_pct=disp_floor_pct,
            jump_high=jump_high,
            jump_low=jump_low,
            jump_min_base=jump_min_base,
            tarifa_estimada=tarifa_estimada,
            duplicates_agg=duplicates_agg,
        )

        meses = sorted({it.get("mes_ref", "") for it in items if it.get("mes_ref")})
        group = {"uc": uc, "concessionaria": conc, "meses_ref": meses, **stat, "anomalias": stat["anomalias"][:max_anom]}
        dups_local = stat.get("duplicidades_local") or []
        if dups_local:
            duplicidades_global.extend(dups_local)
        groups.append(group)
        anomalias_flat.extend(group["anomalias"])

    anomalias_flat.sort(key=lambda x: float(x.get("dif_abs") or 0), reverse=True)

    dup_counter: Counter = Counter()
    for it in prepared:
        key = (it.get("uc", ""), it.get("concessionaria", ""), it.get("mes_ref", ""))
        if any(key):
            dup_counter[key] += 1

    duplicidades_count = []
    for (uc, conc, mes), qtd in dup_counter.items():
        if qtd <= 1:
            continue
        vals = [x.get("value") for x in grouped_raw.get((uc, conc), []) if x.get("mes_ref", "") == mes and x.get("value") is not None]
        media = mean([float(v) for v in vals]) if vals else 0.0
        duplicidades_count.append({"uc": uc, "concessionaria": conc, "mes_ref": mes, "qtd": qtd, "media_kwh": float(media), "alerta": "DUPLICIDADE/AGREGACAO"})

    duplicidades_all = duplicidades_global + duplicidades_count
    duplicidades_all.sort(key=lambda x: (x.get("qtd") or 0), reverse=True)

    result = {
        "ok": True,
        "generated_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "summary": {
            "total_registros": len(prepared),
            "total_grupos": len(groups),
            "total_anomalias": len(anomalias_flat),
            "placeholder_values": sorted(list(placeholder_values)),
        },
        "params": {
            "unit_col": args.unit_col,
            "concessionaria_col": args.concessionaria_col,
            "value_col": args.value_col,
            "days_col": args.days_col,
            "tarifa_col": tarifa_col or None,
            "mes_col": mes_col,
            "min_base": min_base,
            "score_threshold": score_threshold,
            "pct_high": pct_high,
            "pct_low": pct_low,
            "k": k,
            "method": method,
            "baseline_mode": baseline_mode,
            "lookback": lookback,
            "normalize_by_days": normalize_by_days,
            "days_coverage_threshold": days_cov,
            "min_disp": min_disp,
            "disp_floor_pct": disp_floor_pct,
            "leitura_estimada_n": leitura_estimada_n,
            "jump_high": jump_high,
            "jump_low": jump_low,
            "jump_min_base": jump_min_base,
            "duplicates_agg": duplicates_agg,
            "include_series": include_series,
            "series_limit": series_limit,
            "tarifa_estimada": tarifa_estimada,
        },
        "groups": groups,
        "anomalias": anomalias_flat[:max_anom],
        "duplicidades": duplicidades_all[:max_anom],
    }

    safe = make_json_safe(result)
    if args.pretty:
        print(json.dumps(safe, ensure_ascii=False, indent=2))
    else:
        print(json.dumps(safe, ensure_ascii=False))


if __name__ == "__main__":
    main()
